import pdfplumber
import openpyxl
import pandas as pd
from docx import Document


def read_pdf(file_path):
    print("Reading PDF!")
    text = ""
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            text += page.extract_text()
    return text


def read_xlsx(file_path):
    print("Reading XLSX!")
    workbook = openpyxl.load_workbook(file_path)
    sheet = workbook.active
    data = []
    for row in sheet.iter_rows(values_only=True):
        data.append(row)
    print("Data: ", data)
    return data


def read_csv(file_path):
    df = pd.read_csv(file_path)
    return df.to_string()


def read_docx(file_path):
    doc = Document(file_path)
    text = ""
    for para in doc.paragraphs:
        text += para.text + "\n"
    return text


def determine_file_type(file_path):
    print("Determining File Type!")
    if file_path.endswith('.pdf'):
        return 'pdf'
    elif file_path.endswith('.xlsx'):
        print("xlsx file")
        return 'xlsx'
    elif file_path.endswith('.csv'):
        return 'csv'
    elif file_path.endswith('.docx'):
        return 'docx'
    else:
        raise ValueError("Unsupported file type")


def read_file(file_path):
    print("Reached File Reader!")
    file_type = determine_file_type(file_path)
    if file_type == 'pdf':
        return read_pdf(file_path)
    elif file_type == 'xlsx':
        return read_xlsx(file_path)
    elif file_type == 'csv':
        return read_csv(file_path)
    elif file_type == 'docx':
        return read_docx(file_path)
